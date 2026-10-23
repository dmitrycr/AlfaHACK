"""Справочники секторов и компаний из players_by_sector.xlsx -> config/sectors.yaml, config/aliases.yaml, config/okved_to_sector.json.
Лист «Сектора»: точные названия (копируются в Event.sectors, Profile.sector, CounterpartyShare.sector символ в символ) и ОКВЭД.
Лист «Игроки по секторам»: компании и тикеры. ОКВЭД в формате «B.06, C.19, H.49.5, C.10-11, S»:
буква раздела, класс или диапазон классов, группа (49.5). Сначала сопоставление по группе, затем по классу, остальное = «Прочее».
Алиасы, добавленные вручную в прежнем aliases.yaml, переносятся по совпадению имени; факторы sectors.yaml сохраняются.
Использование: python scripts/make_sectors_from_players.py [--xlsx FILE] [--config-dir DIR] [--data-dir DIR]"""
import _bootstrap  # noqa: F401
import argparse
import csv
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
import pandas as pd
import yaml

if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    sys.stdout.reconfigure(encoding="utf-8")

from ml.config import CONFIG_DIR, ML_ROOT
from ml.refdata import normalize_okved, section_of_class

DEFAULT_XLSX = Path(r"C:\Users\fury\projects\AlfaHACK_out\players_by_sector.xlsx")
OTHER = "Прочее"      # TODO(align): служебный сектор для ОКВЭД вне players_by_sector; в Event.sectors не попадает, нужно согласовать с командой

# Ключевые слова (начала слов в нижнем регистре; «re:...» - регулярное выражение). Основа - прежний sectors.yaml, разнесённый по новым секторам.
KEYWORDS: dict[str, list[str]] = {
    "Нефть и газ": ["нефт", "газ", "нпз", "лукойл", "роснефт", "газпром", "новатэк", "баррел"],
    "Горнодобыча (уголь, драгметаллы, алмазы)": ["горнодобыв", "руд", "уголь", "угольн", "золотодобыв", "алмаз", "драгметалл",
                                                  "никел", "медь", "алроса", "полюс", "норникел"],
    "Металлургия": ["металлург", "сталь", "прокат", "чугун", "никел", "медь", "алюмини", "норникел", "северсталь", "нлмк", "ммк",
                    "русал", "металлопрокат", "арматур"],
    "Химия, удобрения, нефтехимия": ["хими", "удобрен", "азотн", "калийн", "фосфор", "нефтехими", "полимер", "полиэтилен", "аммиак"],
    "Энергетика и коммунальные услуги": ["энергет", "электроэнерг", "электростанц", "генераци", "россет", "русгидро", "интер рао", "аэс",
                                         "коммунальн", "жкх", "водоснабжен", "теплоснабжен"],
    "Банки, страхование, финансы": ["банк", "страхов", "цб", "центробанк", "re:центральн\\w+ банк", "облигац", "re:ключев\\w+ ставк", "вклад"],
    "Розничная торговля и e-commerce": ["рознич", "магазин", "ритейл", "торговые сети", "торговых сет", "маркетплейс", "магнит", "ленту"],
    "Оптовая торговля и дистрибуция": ["оптов", "дистрибь", "дистрибу", "торговых компани"],
    "Пищевая промышленность и агро": ["агро", "аграри", "зерн", "урожай", "пшениц", "сельск", "мясо", "молоч", "пищев", "продовольств",
                                      "подсолнеч", "засух"],
    "Машиностроение, автопром, ОПК": ["автомобил", "автопром", "автодилер", "легков", "автоваз", "камаз", "автосалон", "утильсбор",
                                      "машиностро", "станкостро", "оборудован", "промышленн", "оборонн", "вагоностро", "судостро", "станок"],
    "Строительство и девелопмент": ["строит", "застройщик", "новостройк", "ипотек", "стройматериал", "цемент", "девелопер", "жилых проект",
                                    "недвижим", "квартир", "аренд", "жилая", "жилую", "торговые площад", "торговых площад"],
    "Транспорт и логистика": ["транспорт", "логистик", "перевоз", "грузоперевоз", "ржд", "железнодорож", "порт", "авиакомпан", "авиаперевоз",
                              "судоходств", "контейнер", "пассажиропоток"],
    "IT и телеком": ["ит-компан", "программн", "телеком", "мобильн", "облачн", "цод", "яндекс", "мтс", "ростелеком", "кибер"],
    "Фарма и здравоохранение": ["здравоохран", "больниц", "поликлиник", "лекарств", "фармац", "медицин", "клиник", "омс"],
    "Лесная и целлюлозно-бумажная промышленность": ["лесн", "древесин", "лесозаготов", "целлюлоз", "бумаг", "пиломатериал", "фанер", "лесопромышлен"],
    "Общепит и потребительские услуги": ["общепит", "ресторан", "кафе", "фастфуд", "консалтинг", "клининг", "гостиниц", "туризм"],
    OTHER: [],
}

_TOKEN = re.compile(r"^([A-U])(?:\.(\d{2})(?:-(\d{2})|\.(\d))?)?$")      # S | B.06 | C.10-11 | H.49.5


def parse_okved_cell(cell: str) -> list[dict]:
    """'B.06, C.19, H.49.5' -> [{"kind": "class"|"group"|"section", "letter", "classes": [...], "group": "49.5"}]."""
    out = []
    for tok in (t.strip() for t in str(cell).split(",") if t.strip()):
        m = _TOKEN.match(tok)
        if not m:
            raise SystemExit(f"не разобран фрагмент ОКВЭД {tok!r} в {cell!r}")
        letter, c1, c2, grp = m.groups()
        if c1 is None:
            classes = [f"{n:02d}" for n in range(1, 100) if section_of_class(f"{n:02d}") == letter]
            out.append({"kind": "section", "letter": letter, "classes": classes, "token": tok})
        elif grp is not None:
            out.append({"kind": "group", "letter": letter, "classes": [c1], "group": f"{c1}.{grp}", "token": tok})
        else:
            lo, hi = int(c1), int(c2 or c1)
            out.append({"kind": "class", "letter": letter, "classes": [f"{n:02d}" for n in range(lo, hi + 1)], "token": tok})
    return out


def read_workbook(path: Path) -> tuple[list[tuple[str, str]], list[dict]]:
    wb = openpyxl.load_workbook(path, data_only=True)
    sectors = [(r[1], r[2]) for r in wb["Сектора"].iter_rows(min_row=2, values_only=True) if isinstance(r[0], int)]
    players = [{"sector": r[0], "okved": r[1], "name": str(r[2]).strip(), "ticker": (str(r[3]).strip() if r[3] else None)}
               for r in wb["Игроки по секторам"].iter_rows(min_row=2, values_only=True) if r[2]]
    return sectors, players


def build_okved_map(sectors: list[tuple[str, str]]) -> tuple[dict, dict, list[str]]:
    """-> (classes, groups, неоднозначные случаи). Группа сильнее класса, при споре двух секторов за класс побеждает первый по листу."""
    classes: dict[str, str] = {}
    groups: dict[str, str] = {}
    notes: list[str] = []
    for name, cell in sectors:
        for t in parse_okved_cell(cell):
            if t["kind"] == "section":
                notes.append(f"{name}: «{t['token']}» = весь раздел {t['letter']}, классы {t['classes'][0]}-{t['classes'][-1]}")
            for cls in t["classes"]:
                if section_of_class(cls) != t["letter"]:
                    notes.append(f"{name}: «{t['token']}» - класс {cls} относится к разделу {section_of_class(cls)}, не {t['letter']}")
            if t["kind"] == "group":
                if t["group"] in groups and groups[t["group"]] != name:
                    notes.append(f"группа {t['group']} заявлена секторами «{groups[t['group']]}» и «{name}», оставлен первый")
                else:
                    groups[t["group"]] = name
                continue
            for cls in t["classes"]:
                if cls in classes and classes[cls] != name:
                    notes.append(f"класс {cls} заявлен секторами «{classes[cls]}» и «{name}», оставлен первый")
                else:
                    classes[cls] = name
    for g, name in groups.items():
        owner = classes.get(g[:2])
        if owner and owner != name:
            notes.append(f"группа {g} -> «{name}» внутри класса {g[:2]} -> «{owner}»: группа имеет приоритет")
        elif not owner:
            notes.append(f"группа {g} -> «{name}»: класс {g[:2]} целиком ни за кем не закреплён, остальные группы класса -> «{OTHER}»")
    return dict(sorted(classes.items())), dict(sorted(groups.items())), notes


def sector_aliases(names: list[str], rename_csv: Path) -> tuple[dict[str, list[str]], list[str]]:
    """Алиасы названий секторов для приведения ответа модели к каноничным: части названия («Машиностроение, автопром»),
    название без скобок и прежние названия из sector_rename.csv. Алиас, который подходит двум секторам, не добавляется."""
    cand: dict[str, set[str]] = defaultdict(set)         # casefold-алиас -> сектора, которым он подходит
    shown: dict[str, str] = {}
    for n in names:
        base = re.sub(r"\s*\(.*?\)", "", n).strip()
        parts = [p.strip() for p in re.split(r",| и ", base) if p.strip()]
        commas = [p.strip() for p in base.split(",")]
        variants = {base, *parts, *(", ".join(commas[:k]) for k in range(1, len(commas)))}
        for v in variants:
            if v.casefold() != n.casefold() and len(v) >= 2:
                cand[v.casefold()].add(n)
                shown.setdefault(v.casefold(), v)
    notes = []
    if rename_csv.exists():
        old: dict[str, set[str]] = defaultdict(set)
        with open(rename_csv, encoding="utf-8", newline="") as f:
            for r in csv.DictReader(f):
                if r["old"] and r["old"] != r["new"] and r["new"] in names:
                    old[r["old"]].add(r["new"])
        for o, targets in old.items():
            if len(targets) == 1:
                cand[o.casefold()].update(targets)
                shown.setdefault(o.casefold(), o)
            else:
                notes.append(f"прежнее название «{o}» подходит нескольким секторам ({', '.join(sorted(targets))}): алиас не добавлен")
    out: dict[str, list[str]] = {n: [] for n in names}
    canon = {n.casefold() for n in names}
    for key, owners in cand.items():
        if key in canon:
            continue
        if len(owners) == 1:
            out[next(iter(owners))].append(shown[key])
        else:
            notes.append(f"алиас «{shown[key]}» подходит нескольким секторам ({', '.join(sorted(owners))}): не добавлен")
    return {n: sorted(v, key=str.casefold) for n, v in out.items()}, notes


def old_aliases(path: Path) -> dict[str, list[str]]:
    if not path.exists():
        return {}
    raw = (yaml.safe_load(path.read_text(encoding="utf-8")) or {}).get("companies", {})
    return {n: list(i.get("aliases", [])) for n, i in raw.items()}


def q(s: str) -> str:
    return json.dumps(s, ensure_ascii=False)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", default=str(DEFAULT_XLSX))
    ap.add_argument("--config-dir", default=str(CONFIG_DIR))
    ap.add_argument("--data-dir", default=str(ML_ROOT / "data" / "real"), help="для отчёта: куда попадают классы ОКВЭД реальных данных")
    a = ap.parse_args()
    cfg_dir = Path(a.config_dir)
    sectors, players = read_workbook(Path(a.xlsx))
    names = [n for n, _ in sectors]
    if len(names) != 16 or len(set(names)) != 16:
        raise SystemExit(f"ожидалось 16 секторов, найдено {len(names)}")
    if missing := set(names) - set(KEYWORDS):
        raise SystemExit(f"нет ключевых слов для секторов: {missing}")
    bad = {p["sector"] for p in players} - set(names)
    if bad:
        raise SystemExit(f"в листе игроков сектора вне листа «Сектора»: {bad}")
    print(f"секторов: {len(names)}; компаний: {len(players)}")

    # ---- okved_to_sector.json ----
    classes, groups, notes = build_okved_map(sectors)
    mismatch = {(p["sector"], p["okved"]) for p in players} - set(sectors)
    notes += [f"ОКВЭД игрока на листе игроков не совпадает с листом «Сектора»: {s} / {o}" for s, o in sorted(mismatch)]
    (cfg_dir / "okved_to_sector.json").write_text(json.dumps({
        "_todo": f"TODO(align): «{OTHER}» - служебный сектор для ОКВЭД вне players_by_sector; сгенерировано scripts/make_sectors_from_players.py",
        "default": OTHER, "groups": groups, "classes": classes}, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")

    # ---- sectors.yaml (факторы сохраняются из прежнего файла) ----
    old = yaml.safe_load((cfg_dir / "sectors.yaml").read_text(encoding="utf-8"))
    factors = old["factors"]
    prev_enabled = {n: bool(v.get("enabled", True)) for n, v in old["sectors"].items() if isinstance(v, dict)}
    aliases, alias_notes = sector_aliases(names, cfg_dir / "sector_rename.csv")
    notes += alias_notes
    lines = ["# Сгенерировано scripts/make_sectors_from_players.py из players_by_sector.xlsx (лист «Сектора»): названия секторов точные,",
             "# копируются в Event.sectors, Profile.sector, CounterpartyShare.sector символ в символ. Факторы и флаги enabled сохраняются.",
             "# enabled: false - сектор не попадает в промпт тегирования, матрицу и экспорт; профиль клиента строится, но с пометкой out_of_scope.",
             "# keywords: начала слов в нижнем регистре; «re:...» - регулярное выражение. aliases: другие написания названия (ответ модели -> каноничное).",
             f"# TODO(align): «{OTHER}» - служебный сектор (ОКВЭД вне players_by_sector), выключен.",
             "# Включить сектор: docs/add_sector.md (enabled: true, строки матрицы и сырья, алиасы) - код менять не нужно.",
             "sectors:"]
    for n in [*names, OTHER]:
        lines += [f"  {q(n)}:", f"    enabled: {'true' if prev_enabled.get(n, n != OTHER) else 'false'}",
                  f"    keywords: [{', '.join(q(w) for w in KEYWORDS[n])}]",
                  f"    aliases: [{', '.join(q(w) for w in aliases.get(n, []))}]"]
    lines.append("factors:")
    for n, words in factors.items():
        lines.append(f"  {q(n)}: [{', '.join(q(w) for w in words)}]")
    (cfg_dir / "sectors.yaml").write_text("\n".join(lines) + "\n", encoding="utf-8")

    # ---- aliases.yaml ----
    prev = old_aliases(cfg_dir / "aliases.yaml")
    prev_fold = {n.casefold(): n for n in prev} | {al.casefold(): n for n, als in prev.items() for al in als}
    carried, used_old = 0, set()
    lines = ["# Сгенерировано scripts/make_sectors_from_players.py из players_by_sector.xlsx (лист «Игроки по секторам»): имя, тикер, сектор.",
             "# Алиасы из прежнего файла переносятся по совпадению имени. Алиасы до 4 символов ищутся только в точном написании",
             "# (правило в tagging/aliases.py): «пик» - не компания ПИК.", "companies:"]
    for p in players:
        als = [p["name"]] + ([p["ticker"]] if p["ticker"] else [])
        old_name = prev_fold.get(p["name"].casefold()) or (prev_fold.get(p["ticker"].casefold()) if p["ticker"] else None)
        if old_name:
            used_old.add(old_name)
            extra = [x for x in prev[old_name] + ([old_name] if old_name != p["name"] else []) if x not in als]
            carried += bool(extra)
            als += extra
        lines.append(f"  {q(p['name'])}: {{sector: {q(p['sector'])}, aliases: [{', '.join(q(x) for x in dict.fromkeys(als))}]}}")
    (cfg_dir / "aliases.yaml").write_text("\n".join(lines) + "\n", encoding="utf-8")
    notes += [f"компания из прежнего aliases.yaml не найдена в xlsx и удалена: {n}" for n in prev if n not in used_old]

    # ---- отчёт ----
    print(f"classes: {len(classes)}, groups: {len(groups)}; алиасов перенесено для компаний: {carried}")
    print("\nнеоднозначные случаи:")
    for n in notes:
        print(" -", n)
    data_csv = Path(a.data_dir) / "df_Uid_data_hack.csv"
    if data_csv.exists():
        codes = pd.read_csv(data_csv, sep=";", usecols=["MAINOKVED_CODE"], dtype=str)["MAINOKVED_CODE"].dropna()
        cnt: dict[str, Counter] = defaultdict(Counter)
        for code in codes:
            o = normalize_okved(code)
            sec = groups.get(o["group"]) or classes.get(o["class"]) or OTHER
            cnt[sec][o["class"]] += 1
        print(f"\nреальные данные ({len(codes)} клиентов): по секторам")
        for s, c in sorted(cnt.items(), key=lambda kv: -sum(kv[1].values())):
            print(f"  {s}: {sum(c.values())}")
        print(f"классы реальных данных в «{OTHER}»: " + ", ".join(f"{k} ({v})" for k, v in sorted(cnt[OTHER].items())))
    return 0


if __name__ == "__main__":
    sys.exit(main())
